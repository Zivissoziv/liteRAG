"""生成用于端到端验证的样例文档：docx / xlsx / pdf"""
import os
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "samples")
os.makedirs(OUT, exist_ok=True)

# ---------------- DOCX ----------------
doc = Document()
doc.add_heading("LiteRAG 产品需求说明", level=1)
doc.add_paragraph(
    "本文档描述 LiteRAG 知识库系统的产品定位、功能范围与验收标准，用于指导研发与测试工作。"
)

doc.add_heading("一、产品定位", level=2)
doc.add_paragraph(
    "LiteRAG 面向企业内部知识管理场景，把散落在 Word、PDF、Excel 中的非结构化内容转化为"
    "可被大模型检索引用的知识片段。系统强调部署轻量、接入简单，单机即可完成从文档上传到语义检索的全过程。"
)

doc.add_heading("二、功能需求", level=2)
doc.add_paragraph("1. 文档管理：支持 docx、pdf、xlsx、pptx、markdown 等格式上传，自动解析并分块入库。")
doc.add_paragraph("2. 知识库隔离：不同业务域使用独立知识库，各自绑定独立的向量集合与 Embedding 模型。")
doc.add_paragraph("3. 检索能力：支持按相似度阈值过滤、按知识库范围过滤、按文档过滤，返回带得分的原文片段。")
doc.add_paragraph("4. 模型配置：支持任意 OpenAI 兼容接口与本地 Ollama，可在线测试连通性与向量维度。")
doc.add_paragraph("5. MCP 服务：通过 Model Context Protocol 向外部 AI 客户端提供知识检索工具。")

doc.add_heading("三、性能指标", level=2)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "指标"
hdr[1].text = "目标值"
hdr[2].text = "说明"
for row in [
    ("单文件解析耗时", "小于 10 秒", "10MB 以内的 docx 文档"),
    ("检索响应时间", "小于 200 毫秒", "单知识库 10 万片段以内"),
    ("并发上传", "20 个文件", "单次请求上限"),
]:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value

doc.add_heading("四、验收标准", level=2)
doc.add_paragraph(
    "上传一份 20 页的产品手册后，使用自然语言提问，系统应召回至少一条相关片段，"
    "并在回答中给出准确的来源编号。检索结果的相关性由产品经理人工评估，前三条命中率需达到 80% 以上。"
)
doc.save(os.path.join(OUT, "literag-产品需求.docx"))

# ---------------- XLSX ----------------
wb = Workbook()
ws = wb.active
ws.title = "成本明细"
ws.append(["项目", "单价(元)", "数量", "小计(元)", "负责人"])
rows = [
    ("云服务器 4C8G", 480, 2, 960, "张工"),
    ("Qdrant 向量库节点", 120, 3, 360, "李工"),
    ("向量化调用", 0.0007, 1200000, 840, "王工"),
    ("对象存储", 60, 2, 120, "张工"),
    ("域名与证书", 90, 1, 90, "赵工"),
]
for r in rows:
    ws.append(list(r))
ws.append(["合计", "", "", "=SUM(D2:D6)", ""])

ws2 = wb.create_sheet("里程碑")
ws2.append(["阶段", "开始日期", "结束日期", "交付物"])
for r in [
    ("需求评审", "2026-03-02", "2026-03-06", "需求规格说明书"),
    ("后端开发", "2026-03-09", "2026-03-27", "Spring Boot 服务"),
    ("前端开发", "2026-03-16", "2026-03-31", "Vue3 管理台"),
    ("联调测试", "2026-04-01", "2026-04-10", "测试报告"),
]:
    ws2.append(list(r))
wb.save(os.path.join(OUT, "literag-项目成本.xlsx"))

# ---------------- PDF ----------------
pdf_path = os.path.join(OUT, "literag-deployment-guide.pdf")
c = canvas.Canvas(pdf_path, pagesize=A4)
width, height = A4
lines = [
    ("LiteRAG Deployment Guide", 20),
    ("", 12),
    ("1. Prerequisites", 14),
    ("Java 17 or later, Node.js 18 or later, Docker (optional).", 11),
    ("Qdrant runs as a single binary on port 6333 (or via docker compose).", 11),
    ("", 12),
    ("2. Start the backend", 14),
    ("Run: java -jar literag-server.jar --server.port=8080", 11),
    ("The metadata database is an embedded SQLite file at ./data/db.", 11),
    ("", 12),
    ("3. Start the frontend", 14),
    ("Run: npm install && npm run dev inside the web directory.", 11),
    ("The dev server proxies /api and /mcp to port 8080.", 11),
    ("", 12),
    ("4. Configure an embedding model", 14),
    ("Open the Models page, add an OpenAI compatible endpoint", 11),
    ("such as BAAI/bge-m3, test the connection and set it as default.", 11),
    ("", 12),
    ("5. Connect an MCP client", 14),
    ("Point the client at http://localhost:8080/mcp using the", 11),
    ("Streamable HTTP transport, then call search_knowledge.", 11),
]
y = height - 70
for text, size in lines:
    c.setFont("Helvetica", size)
    c.drawString(70, y, text)
    y -= size + 8
c.showPage()
c.save()

print("samples written to", OUT)
print(os.listdir(OUT))
