"""生成一份用于验证 pptx 版式预览的样例演示文稿。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 封面
s1 = prs.slides.add_slide(prs.slide_layouts[6])
bg = s1.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = RGBColor(0x1F, 0x2D, 0x5C)
bg.line.fill.background()

tb = s1.shapes.add_textbox(Inches(1), Inches(2.4), Inches(11.3), Inches(1.6))
p = tb.text_frame.paragraphs[0]
p.text = "LiteRAG 知识库方案"
p.alignment = PP_ALIGN.LEFT
p.runs[0].font.size = Pt(44)
p.runs[0].font.bold = True
p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

sub = s1.shapes.add_textbox(Inches(1), Inches(4.0), Inches(11.3), Inches(0.8))
p2 = sub.text_frame.paragraphs[0]
p2.text = "文档解析 分块 向量检索 MCP 服务"
p2.runs[0].font.size = Pt(20)
p2.runs[0].font.color.rgb = RGBColor(0xC3, 0xD0, 0xFF)

# 内容页
s2 = prs.slides.add_slide(prs.slide_layouts[6])
title = s2.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.9))
tp = title.text_frame.paragraphs[0]
tp.text = "系统能力"
tp.runs[0].font.size = Pt(32)
tp.runs[0].font.bold = True
tp.runs[0].font.color.rgb = RGBColor(0x20, 0x20, 0x20)

body = s2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(6.0), Inches(4.4))
tf = body.text_frame
tf.word_wrap = True
items = [
    "支持 docx、pdf、xlsx、pptx 上传与解析",
    "段落感知分块，默认 800 字、重叠 120 字",
    "自定义 Embedding：OpenAI 兼容、Ollama、离线",
    "MCP 提供检索与问答工具，可接入任意 AI 客户端",
]
for i, text in enumerate(items):
    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    para.text = "• " + text
    para.font.size = Pt(18)
    para.space_after = Pt(10)
    para.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

# 表格
s3 = prs.slides.add_slide(prs.slide_layouts[6])
t_title = s3.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.9))
q = t_title.text_frame.paragraphs[0]
q.text = "检索性能"
q.runs[0].font.size = Pt(32)
q.runs[0].font.bold = True

rows, cols = 4, 3
table = s3.shapes.add_table(rows, cols, Inches(0.8), Inches(1.8), Inches(9.5), Inches(2.6)).table
data = [
    ("知识库规模", "平均耗时", "命中率"),
    ("1 万片段", "68 ms", "92%"),
    ("10 万片段", "143 ms", "89%"),
    ("100 万片段", "310 ms", "85%"),
]
for r in range(rows):
    for c in range(cols):
        cell = table.cell(r, c)
        cell.text = data[r][c]
        cell.text_frame.paragraphs[0].font.size = Pt(16)

prs.save("scripts/samples/literag-architecture.pptx")
print("saved")
